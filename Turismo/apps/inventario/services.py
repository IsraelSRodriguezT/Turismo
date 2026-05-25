import csv
import io
from typing import List, Dict, Any
from django.core.exceptions import ValidationError
from .models import Recurso, TipoRecurso

def crear_recurso(titulo: str, descripcion: str, url: str, tipo_recurso: str) -> Recurso:
	tipo_norm = (tipo_recurso or '').strip().upper()
	if tipo_norm not in [t.value for t in TipoRecurso]:
		raise ValidationError(f"Tipo de recurso no valido: {tipo_recurso}")
	recurso = Recurso.objects.create(
		titulo=titulo,
		descripcion=descripcion or '',
		url=url,
		tipo_recurso=tipo_norm,
	)
	return recurso

def listar_recursos(tipo_recurso: str | None = None):
	qs = Recurso.objects.all()
	if tipo_recurso:
		qs = qs.filter(tipo_recurso=tipo_recurso)
	return qs

def _validate_csv_headers(headers: List[str]) -> List[str]:
	required = ['titulo', 'descripcion', 'url', 'tipo_recurso']
	seen = [ (x or '').strip().lower() for x in headers ]
	missing = [h for h in required if h not in seen]
	return missing

def _es_duplicado(url: str | None, titulo: str | None) -> bool:
	if url:
		if Recurso.objects.filter(url=url).exists():
			return True
	if titulo:
		if Recurso.objects.filter(titulo__iexact=titulo).exists():
			return True
	return False

def importar_desde_excel(file_obj) -> Dict[str, Any]:
	name = getattr(file_obj, 'name', '')
	filename = name.lower()
	results: Dict[str, Any] = {'created': 0, 'skipped': [], 'errors': []}

	# Manejar CSV
	if filename.endswith('.csv') or getattr(file_obj, 'content_type', '').endswith('csv'):
		data = file_obj.read()
		if isinstance(data, bytes):
			data = data.decode('utf-8')
		reader = csv.DictReader(io.StringIO(data))
		missing = _validate_csv_headers(reader.fieldnames or [])
		if missing:
			raise ValidationError(f'Cabeceras faltantes en CSV: {missing}')
		for idx, row in enumerate(reader, start=1):
			try:
				titulo_val = row.get('titulo', '').strip()
				url_val = row.get('url', '').strip()
				if _es_duplicado(url_val, titulo_val):
					results['skipped'].append({'row': idx, 'reason': 'duplicate', 'titulo': titulo_val, 'url': url_val})
				else:
					crear_recurso(
						titulo=titulo_val,
						descripcion=row.get('descripcion', '').strip(),
						url=url_val,
						tipo_recurso=row.get('tipo_recurso', '').strip(),
					)
					results['created'] += 1
			except Exception as e:
				results['errors'].append({'row': idx, 'error': str(e)})
		return results

	# Manejar XLSX
	if filename.endswith('.xlsx') or filename.endswith('.xls'):
		try:
			import openpyxl
		except Exception:
			raise ValidationError('Para importar XLSX se requiere openpyxl: pip install openpyxl')

		try:
			file_obj.seek(0)
		except Exception:
			pass

		wb = openpyxl.load_workbook(file_obj, read_only=True)
		ws = wb.active

		rows_iter = ws.iter_rows(values_only=True)
		try:
			first = next(rows_iter)
		except StopIteration:
			raise ValidationError('Archivo XLSX vacío')

		headers = [str(h).strip().lower() for h in first]
		missing = _validate_csv_headers(headers)
		if missing:
			raise ValidationError(f'Cabeceras faltantes en XLSX: {missing}')

		for sheet_row_idx, row in enumerate(rows_iter, start=2):
			row_dict = dict(zip(headers, [r if r is not None else '' for r in row]))
			try:
				titulo_val = str(row_dict.get('titulo', '')).strip()
				url_val = str(row_dict.get('url', '')).strip()
				if _es_duplicado(url_val, titulo_val):
					results['skipped'].append({'row': sheet_row_idx, 'reason': 'duplicate', 'titulo': titulo_val, 'url': url_val})
				else:
					crear_recurso(
						titulo=titulo_val,
						descripcion=str(row_dict.get('descripcion', '')).strip(),
						url=url_val,
						tipo_recurso=str(row_dict.get('tipo_recurso', '')).strip(),
					)
					results['created'] += 1
			except Exception as e:
				results['errors'].append({'row': sheet_row_idx, 'error': str(e)})
		return results

	raise ValidationError('Formato de archivo no soportado. Use CSV o XLSX.')

def validar_estructura_archivo(file_obj) -> Dict[str, Any]:
	name = getattr(file_obj, 'name', '')
	filename = name.lower()

	# CSV
	if filename.endswith('.csv') or getattr(file_obj, 'content_type', '').endswith('csv'):
		data = file_obj.read()
		if isinstance(data, bytes):
			data = data.decode('utf-8')
		reader = csv.DictReader(io.StringIO(data))
		headers = reader.fieldnames or []
		missing = _validate_csv_headers(headers)
		# rewind if possible
		try:
			file_obj.seek(0)
		except Exception:
			pass
		return {'ok': not bool(missing), 'missing': missing, 'headers': headers}

	# XLSX
	if filename.endswith('.xlsx') or filename.endswith('.xls'):
		try:
			import openpyxl
		except Exception:
			return {'ok': False, 'missing': ['openpyxl'], 'headers': []}
		# Asegurar file-like en inicio antes de leer
		try:
			file_obj.seek(0)
		except Exception:
			pass
		wb = openpyxl.load_workbook(file_obj, read_only=True)
		ws = wb.active
		rows_iter = ws.iter_rows(values_only=True)
		try:
			first = next(rows_iter)
		except StopIteration:
			return {'ok': False, 'missing': ['empty'], 'headers': []}
		headers = [str(h).strip().lower() for h in first]
		missing = _validate_csv_headers(headers)
		try:
			file_obj.seek(0)
		except Exception:
			pass
		return {'ok': not bool(missing), 'missing': missing, 'headers': headers}

	return {'ok': False, 'missing': ['unsupported_format'], 'headers': []}

